"""Build the 29 October handbook from its canonical prose, preserving literal IDs.

Run with the bundled Codex Python. Never imports or executes attached source code.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
from zipfile import ZipFile
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[3]
BASE=ROOT/'docs/upgrade/2026-10-29'
DEST=BASE/'exports/Life_Code_OS_Master_Dev_Handoff_2026-10-29.docx'
SOURCES=[
 ('Tổng kế hoạch hợp nhất','MASTER_PLAN.md'),
 ('Bàn giao backend và công thức','backend/BACKEND_DEV_HANDOFF.md'),
 ('Sản phẩm và giao diện đa thiết bị','product/PRODUCT_UX_REPORT_HANDOFF.md'),
 ('AI báo cáo và Academy','product/AI_REPORT_ACADEMY_HANDOFF.md'),
 ('Điều hành và nhận việc DEV','execution/PROGRAM_EXECUTION_PLAN.md'),
 ('Đối chiếu API và tra cứu hợp đồng','backend/API_CONTRACT_RECONCILIATION.md'),
]
BLACK=RGBColor(0,0,0)

def strip_markup(text):
    # Remove Markdown presentation only. Do not sanitize punctuation, IDs or dates.
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:m[1]+' ('+m[2]+')',text)
    return text.replace('**','').replace('`','').strip()

def xml(tag,**attrs):
    el=OxmlElement('w:'+tag)
    for k,v in attrs.items():el.set(qn('w:'+k),str(v))
    return el

def set_font(run,size=None,bold=None):
    run.font.name='Arial';run.font.color.rgb=BLACK
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    rpr=run._r.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=xml('rFonts');rpr.insert(0,fonts)
    for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Arial')

def bookmark(p,name,num):
    p._p.insert(0,xml('bookmarkStart',id=num,name=name))
    p._p.append(xml('bookmarkEnd',id=num))

class Book:
    def __init__(self):
        self.doc=Document();self.headings=[];self.toc_rows=[];self.source_receipts=[]
        s=self.doc.sections[0]
        s.page_width=Inches(8.5);s.page_height=Inches(11)
        s.top_margin=Inches(.7);s.bottom_margin=Inches(.7);s.left_margin=Inches(.78);s.right_margin=Inches(.78)
        s.header_distance=Inches(.28);s.footer_distance=Inches(.28)
        for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number','TOC 1','TOC 2']:
            try:st=self.doc.styles[name]
            except KeyError:st=self.doc.styles.add_style(name,1)
            st.font.name='Arial';st.font.color.rgb=BLACK
            st.font.size=Pt(11)
            st.paragraph_format.space_after=Pt(6)
            st.paragraph_format.line_spacing=1.12
            st.paragraph_format.widow_control=True
            ppr=st.element.get_or_add_pPr()
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
            if name.startswith('Heading'):
                lev=int(name[-1]);st.font.size=Pt({1:18,2:14,3:11.5}[lev]);st.font.bold=True
                st.paragraph_format.space_before=Pt(14 if lev>1 else 0)
                st.paragraph_format.keep_with_next=True
            if name in ['Title','Subtitle']:
                st.font.size=Pt(28 if name=='Title' else 19)
                st.font.bold=name=='Title'
            # Explicit black with no themeColor residue, underline or paragraph rule.
            rpr=st.element.get_or_add_rPr()
            fonts=rpr.find(qn('w:rFonts'))
            if fonts is not None:
                for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:fonts.attrib.pop(qn('w:'+attr),None)
                for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Arial')
            color=rpr.find(qn('w:color'))
            if color is not None:
                for k in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+k),None)
            for u in list(rpr.findall(qn('w:u'))):rpr.remove(u)
        self.doc.styles['Normal'].paragraph_format.keep_with_next=False
        header=s.header.paragraphs[0]
        header.add_run('LIFE CODE OS  •  BÀN GIAO TỔNG KẾ HOẠCH VÀ DEV')
        for run in header.runs:set_font(run,8)
        foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        set_font(foot.add_run('2026-10-29.r1   |   Trang '),8)
        field=xml('fldSimple',instr='PAGE');r=xml('r');rpr=xml('rPr');rpr.append(xml('sz',val='16'));rpr.append(xml('color',val='000000'));r.append(rpr);r.append(xml('t'));field.append(r);foot._p.append(field)
        self.doc.core_properties.title='Life Code OS Bàn giao tổng kế hoạch và phát triển'
        self.doc.core_properties.subject='2026-10-29.r1 — tổng hợp V1 V2 V3 và local 20/10.r2'
        self.doc.core_properties.author='Life Code OS'
        self.doc.core_properties.keywords='Life Code OS, DEV handoff, 2026-10-29.r1'

    def para(self,text,style=None):
        p=self.doc.add_paragraph(style=style)
        set_font(p.add_run(strip_markup(text)))
        return p

    def heading(self,text,level):
        text=strip_markup(text)
        p=self.doc.add_heading(text,level)
        for run in p.runs:set_font(run)
        name=f'lc_h{len(self.headings)+1:03}'
        bookmark(p,name,len(self.headings)+1)
        self.headings.append({'text':text,'level':level,'bookmark':name})
        return p

    def table(self,lines):
        rows=[[strip_markup(x) for x in line.strip().strip('|').split('|')] for line in lines]
        rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x.strip()) for x in r)]
        if len(rows)<2:return
        n=max(map(len,rows));headers=rows[0]
        # Long narrative records become readable paragraphs; the original registry stays authoritative.
        narrative=(n>=4 and max(sum(map(len,r)) for r in rows[1:])>650)
        narrative=narrative or (n>=3 and ('Gói' in headers[0] or 'Gói' in ' '.join(headers)) and max(sum(map(len,r)) for r in rows[1:])>280)
        narrative=narrative or (n>=3 and 'WP /' in headers[0])
        narrative=narrative or 'Method candidates V3' in headers
        if narrative:
            for row in rows[1:]:
                self.heading(row[0],3)
                for index,value in enumerate(row[1:],1):
                    if value:self.para(headers[index]+': '+value)
            return
        table=self.doc.add_table(rows=0,cols=n);table.autofit=False
        available=6.94
        if n==2: widths=[2.15,4.79]
        elif n==3:widths=[1.35,2.25,3.34]
        elif n==4:widths=[.8,1.7,2.22,2.22]
        else:widths=[available/n]*n
        if headers[0]=='ID mới' and n==4:widths=[1.38,1.42,2.07,2.07]
        grid=table._tbl.tblGrid
        for i,col in enumerate(grid.gridCol_lst):col.set(qn('w:w'),str(round(widths[i]*1440)))
        for index,row in enumerate(rows):
            cells=table.add_row().cells
            if headers[0] in ['ID mới','Registry']:table.rows[-1]._tr.get_or_add_trPr().append(xml('cantSplit'))
            for j,cell in enumerate(cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=row[j] if j<len(row) else ''
                pr=cell._tc.get_or_add_tcPr()
                borders=xml('tcBorders')
                for edge in ['top','left','bottom','right']:borders.append(xml(edge,val='single',sz='4',color='D9D9D9'))
                pr.append(borders);margins=xml('tcMar')
                for edge in ['top','left','bottom','right']:margins.append(xml(edge,w='100',type='dxa'))
                pr.append(margins)
                pr.append(xml('shd',fill='DCE5ED' if index==0 else 'F5F7F9' if index%2==0 else 'FFFFFF'))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.08
                    for run in p.runs:set_font(run,10,index==0)
            if index==0:
                table.rows[0]._tr.get_or_add_trPr().append(xml('tblHeader'))
        gap=self.doc.add_paragraph();gap.paragraph_format.space_after=Pt(2);gap.paragraph_format.space_before=Pt(0)

    def markdown(self,text,skip_title=True,api_summary=False):
        lines=text.splitlines();i=0;code=False
        while i<len(lines):
            line=lines[i].strip();i+=1
            if line.startswith('```'):code=not code;continue
            if not line or line=='---':continue
            if api_summary and line.startswith('## Sổ 99 operation'):
                self.heading('Tra cứu đủ 99 operation trong registry API',2)
                self.para('Sổ chi tiết 99 dòng được giữ đầy đủ trong backend/API_CONTRACT_RECONCILIATION.json và .md. Tra theo operationId để lấy method/path sau chuẩn hóa, quyền nguồn, request/response/error, provenance và blocker. Không chép sổ này thành bảng sáu cột hẹp trong sách; registry là hợp đồng kèm theo để DEV kiểm tra và reviewer đối chiếu.')
                while i<len(lines) and not lines[i].startswith('## Bàn giao schema'):i+=1
                continue
            if line.startswith('|') and not code:
                group=[line]
                while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
                self.table(group);continue
            match=re.match(r'^(#{1,6})\s+(.+)',line)
            if match and not code:
                lev=len(match[1])
                if lev==1 and skip_title:continue
                self.heading(match[2],min(lev,3));continue
            if code:
                p=self.para(line)
                for run in p.runs:run.font.name='Courier New';run.font.size=Pt(9)
            elif re.match(r'^[-*]\s+',line):self.para(line[2:],'List Bullet')
            elif re.match(r'^\d+\.\s+',line):
                p=self.para(line);p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22)
            else:self.para(line)

    def build(self):
        self.doc.add_paragraph('Life Code OS',style='Title')
        self.doc.add_paragraph('Bàn giao tổng kế hoạch và phát triển',style='Subtitle')
        self.para('Phiên bản 2026-10-29.r1  |  Chủ dự án Trần Hà Tâm')
        self.para('Bộ kế hoạch đã hợp nhất V1, V2, V3 và phần làm chặt nội bộ ngày 20/10.r2 để đội ngũ nhận việc DEV theo phạm vi và bằng chứng. Các kế hoạch cũ, ID và nguồn được bảo toàn. Backend, công thức, quyền dữ liệu, demo, báo cáo và trải nghiệm mobile web đến desktop web có hợp đồng bàn giao riêng trong cùng bộ.')
        self.para('Kết luận hiện tại là có thể nhận việc chuẩn bị, đặc tả và phát triển synthetic theo work order. Phát hành vẫn HOLD cho tới khi candidate, nguồn và phương pháp, quyền, staging, phục hồi, thương mại và người duyệt có receipt đúng phạm vi. Tài liệu này không tuyên bố ứng dụng production đã được nâng cấp.')
        self.para('Dành cho Founder, Product, Method, Backend, Frontend, AI, Editorial, QA, Security, Operations, Legal và Finance. Đọc tổng kế hoạch trước, sau đó dùng phần chuyên môn và các registry được dẫn để giao task, kiểm phụ thuộc và thu bằng chứng.')
        self.para('Đối chiếu nguồn và bản máy đọc nằm trong docs/upgrade/2026-10-29/. Các bảng đề xuất không cấp approval, giá, ngân sách, quyền dữ liệu hay quyền deploy.')
        self.doc.add_page_break()
        toc_title=self.doc.add_paragraph('Mục lục',style='Heading 1')
        # Temporarily move the TOC to a separate document position after generating headings.
        toc_anchor=self.doc.add_paragraph('__TOC_ANCHOR__')
        self.doc.add_page_break()
        for chapter,(title,relative) in enumerate(SOURCES,1):
            # Appendices flow after short closing paragraphs instead of creating near-blank leaves.
            if chapter>1 and chapter not in [5,6]:self.doc.add_page_break()
            self.heading(f'Phần {chapter} {title}',1)
            p=BASE/relative;text=p.read_text()
            self.source_receipts.append({'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})
            self.para('Tài liệu chuyên môn kèm theo: '+relative)
            self.markdown(text,api_summary=chapter==6)
        self.heading('Phần 7 Danh mục hợp đồng và cách nhận bàn giao',1)
        self.para('Bản sách là hướng dẫn đọc và nhận việc. Các registry sau giữ toàn bộ bản ghi chi tiết và provenance để công cụ DEV/QA kiểm tra. Không dùng số trang sách làm mẫu số hoàn thành dự án.')
        inventory=[
          ('execution/BACKLOG_UNIFIED.json','296 task nguồn và 20 work order tích hợp; 316 bản ghi, effort không cộng trùng'),
          ('execution/TASK_DEV_CONTRACTS.json','316 hợp đồng nhận việc; owner/reviewer thật, candidate, scope và receipt cần được điền khi dispatch'),
          ('execution/TASK_DEPENDENCY_GRAPH.json','Task, decision và gate được tách kiểu; không đổi approval thành task hoàn tất'),
          ('execution/DECISION_EXECUTION_PACKETS.json','Q1–Q45, giữ nguyên packet cũ và ghi các quyết định còn pending'),
          ('execution/INTEGRATION_WORK_ORDERS.json','INT29-001…INT29-020, scope và acceptance phân rã tích hợp'),
          ('execution/REQUIREMENT_TRACE_UNIFIED.json','407 dòng trace, với grain nguồn/product/screen/method/API riêng'),
          ('backend/OPENAPI_UNIFIED_TARGET.json','99 operation, 68 component schema; codegen chưa sẵn sàng vì 81 operation có blocker'),
          ('backend/API_CONTRACT_RECONCILIATION.json','Đối chiếu từng API, permission, schema, path và blockers'),
          ('backend/FORMULA_METHOD_EXECUTION_MATRIX.json','62 method, namespace passport/formula, entrypoint reference và việc còn thiếu'),
          ('backend/SOURCE_DATA_INTEGRATION_MATRIX.json','Asset/dictionary/table/source IDs, hash, quyền và admission trước seed'),
          ('product/PRODUCT_EXECUTION_CONTRACTS.json','28 sản phẩm; 23 rich contract cũ được giữ nguyên trường'),
          ('product/SCREEN_EXECUTION_CONTRACTS.json','69 ID màn hình; 45 rich contract cũ và 24 đề xuất bổ sung'),
          ('product/NINE_LEVEL_REPORT_CONTRACTS.json','Chín report contract; 54 fixture cũ và 36 fixture bổ sung được yêu cầu'),
          ('evidence/TEST_RECEIPT_BACKEND.json','Lệnh và giới hạn lần kiểm synthetic hiện tại; không thay staging/runtime'),
          ('evidence/READ_COVERAGE_PRODUCT.json','46 tệp đọc trong lane sản phẩm và kiểm tra liên kết hợp đồng'),
          ('evidence/DOCUMENT_QA.json','Receipt cấu trúc, render và kiểm tra từng trang của bản sách này'),
        ]
        self.table(['| Registry | Phạm vi tra cứu |','|---|---|']+['|'+p+'|'+meaning+'|' for p,meaning in inventory])
        self.heading('Nhận việc theo candidate và phạm vi',2)
        self.para('Người nhận task kiểm version và SHA nguồn, chọn capability/cohort, gán người chịu trách nhiệm và reviewer độc lập, xác nhận file allowlist, decision/gate cần thiết, fixture và lệnh kiểm tra. Đầu ra phải có expected/actual, artifact digest, môi trường, thời gian, giới hạn và rollback. Local pass, staging pass và release authorization là các receipt khác nhau.')
        self.para('Giữ các nguồn và kế hoạch cũ tại vị trí bảo toàn. Nếu cần đổi hành vi, tạo amendment có ID và lý do; không sửa ngầm nguyên bản hoặc xóa yêu cầu chưa nghiên cứu đủ. Chưa có bằng chứng cho một capability thì giữ trạng thái pending/blocked cùng owner role và sự kiện cần để mở lại.')
        # Static, linked TOC supports deterministic headless delivery and real page numbers after render.
        for item in self.headings:
            if item['level']>2:continue
            p=self.doc.add_paragraph(style='TOC 1' if item['level']==1 else 'TOC 2')
            p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.left_indent=Inches(0 if item['level']==1 else .18)
            tabs=p.paragraph_format.tab_stops
            from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
            tabs.add_tab_stop(Inches(6.75),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            link=xml('hyperlink',anchor=item['bookmark'],history='1')
            r=xml('r');rp=xml('rPr');rp.append(xml('color',val='000000'));rp.append(xml('sz',val='20'));r.append(rp)
            text=xml('t');text.text=item['text'];r.append(text);link.append(r)
            p._p.append(link)
            set_font(p.add_run('\t'))
            page=xml('r');text=xml('t');text.text='1';page.append(text);p._p.append(page)
            p._p.set(qn('w:rsidR'),'00C0D0E0')
            toc_anchor._p.addprevious(p._p)
            self.toc_rows.append(item)
        toc_anchor._p.getparent().remove(toc_anchor._p)
        DEST.parent.mkdir(exist_ok=True)
        self.doc.save(DEST)
        meta={'version':'2026-10-29.r1','sources':self.source_receipts,'headings':self.headings,'toc_entries':self.toc_rows,'toc_kind':'static internal hyperlinks and bookmarks; page numbers populated from rendered PDF outlines','docx_path':str(DEST),'paragraph_count':len(self.doc.paragraphs),'table_count':len(self.doc.tables)}
        (BASE/'evidence/DOCUMENT_BUILD.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n')
        print(json.dumps({'docx':str(DEST),'headings':len(self.headings),'toc_entries':len(self.toc_rows),'tables':len(self.doc.tables)},ensure_ascii=False))

def fill_toc(pdf):
    from pypdf import PdfReader
    reader=PdfReader(str(pdf));mapping={}
    def walk(items):
        for item in items:
            if isinstance(item,list):walk(item)
            else:
                try:mapping[item.title]=reader.get_destination_page_number(item)+1
                except (AttributeError,ValueError):pass
    walk(reader.outline)
    doc=Document(DEST);filled=[]
    for p in doc.paragraphs:
        if p.style.name not in ['TOC 1','TOC 2']:continue
        title=''.join(p._p.xpath('./w:hyperlink/w:r/w:t/text()'))
        if title not in mapping:raise ValueError('TOC heading absent from rendered PDF outline: '+title)
        texts=p._p.xpath('./w:r/w:t');assert texts
        texts[-1].text=str(mapping[title]);filled.append({'title':title,'page':mapping[title]})
    assert filled
    doc.save(DEST)
    (BASE/'evidence/DOCUMENT_TOC.json').write_text(json.dumps({'entries':filled,'pdf_pages':len(reader.pages)},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'filled_toc_entries':len(filled),'pdf_pages':len(reader.pages)}))

def qa(pdf,render_dir,visual_reviewed):
    pdf=pdf.resolve();render_dir=render_dir.resolve()
    from pypdf import PdfReader
    from lxml import etree
    import pdfplumber
    reader=PdfReader(str(pdf));doc=Document(DEST)
    build=json.loads((BASE/'evidence/DOCUMENT_BUILD.json').read_text())
    toc=json.loads((BASE/'evidence/DOCUMENT_TOC.json').read_text())
    images=sorted(render_dir.glob('page-*.png'),key=lambda p:int(p.stem.split('-')[-1]))
    assert len(images)==len(reader.pages)>0
    assert [int(p.stem.split('-')[-1]) for p in images]==list(range(1,len(images)+1))
    for source in build['sources']:assert hashlib.sha256((ROOT/source['path']).read_bytes()).hexdigest()==source['sha256'],('Source changed after build',source['path'])
    outlines={}
    def walk(items):
        for item in items:
            if isinstance(item,list):walk(item)
            else:
                try:outlines[item.title]=reader.get_destination_page_number(item)+1
                except (AttributeError,ValueError):pass
    walk(reader.outline)
    for entry in toc['entries']:assert outlines.get(entry['title'])==entry['page'],('TOC page mismatch',entry)
    assert all(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for t in doc.tables)
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name]
        assert style.font.color.rgb==BLACK
        assert not style.element.xpath('./w:pPr/w:pBdr')
    assert not any(p._p.xpath('./w:pPr/w:pBdr') for p in doc.paragraphs if p.style.name in ['Title','Subtitle'])
    assert any('PAGE' in p._p.xml for p in doc.sections[0].footer.paragraphs)
    all_doc_text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert '__TOC_ANCHOR__' not in all_doc_text
    pdf_text='\n'.join(p.extract_text() or '' for p in reader.pages)
    literal_checks={token:{'docx':token in all_doc_text,'pdf':token in pdf_text} for token in ['2026-10-29.r1','INT29-001','Q1–Q45','125/125','192/192','81/99','62/62','LC-NUM-ATLAS','V3-UI-058']}
    assert all(all(check.values()) for check in literal_checks.values()),literal_checks
    outer=[];pages=[]
    with pdfplumber.open(pdf) as parsed:
        for index,page in enumerate(parsed.pages,1):
            chars=[c for c in page.chars if c.get('text','').strip()]
            bad=[c for c in chars if c['x0']<24 or c['x1']>page.width-24 or c['top']<14 or c['bottom']>page.height-14]
            if bad:outer.append({'page':index,'characters':len(bad)})
            text=page.extract_text() or ''
            assert len(text)>100,('Unexpected blank/near-blank page',index)
            pages.append({'page':index,'text_characters':len(text),'png':str(images[index-1]),'png_sha256':hashlib.sha256(images[index-1].read_bytes()).hexdigest(),'visual_status':'PASS' if visual_reviewed else 'NOT_CHECKED'})
    assert not outer,outer
    receipt={
       'version':'2026-10-29.r1','scope':'NEW_MASTER_HANDOFF_DOCUMENTS_ONLY_NOT_APPLICATION_OR_RELEASE_QA',
       'status':'PASS' if visual_reviewed else 'STRUCTURAL_PASS_VISUAL_PENDING',
       'docx':{'path':str(DEST.relative_to(ROOT)),'sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),'bytes':DEST.stat().st_size,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},
       'pdf':{'path':str(pdf.relative_to(ROOT)),'sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'bytes':pdf.stat().st_size,'pages':len(reader.pages)},
       'sources':build['sources'],'toc':{'kind':build['toc_kind'],'entries':len(toc['entries']),'all_final_page_numbers_match_pdf_outline':True},
       'structural_checks':{'used_heading_styles_black':True,'title_and_subtitle_no_paragraph_border':True,'all_tables_repeating_headers':True,'footer_PAGE_field':True,'all_pages_rendered':True,'no_characters_outside_24pt_horizontal_14pt_vertical_safety_inset':True,'no_blank_pages':True,'source_hashes_match_build':True},
       'literal_preservation':literal_checks,'visual_review':{'all_pages_individually_reviewed':visual_reviewed,'resolution':'original pixel resolution; every final page PNG, not contact sheets only','pages':pages},
       'runtime':{'python':'/Users/tranhatam/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3','libreoffice':'/Users/tranhatam/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice','renderer':'/Users/tranhatam/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.22136/skills/documents/render_docx.py'},
       'limits':['Document layout and planning consistency only.','No application, staging, production, commercial, legal or method release approval.','PDF is exported from the same DOCX content; static linked TOC is regenerated by the builder after edits.']}
    (BASE/'evidence/DOCUMENT_QA.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'status':receipt['status'],'pages':len(reader.pages),'toc_entries':len(toc['entries']),'tables':len(doc.tables)}))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--fill-toc',type=Path);parser.add_argument('--qa',type=Path);parser.add_argument('--render-dir',type=Path);parser.add_argument('--visual-reviewed',action='store_true');args=parser.parse_args()
    if args.qa:qa(args.qa,args.render_dir,args.visual_reviewed)
    elif args.fill_toc:fill_toc(args.fill_toc)
    else:Book().build()
