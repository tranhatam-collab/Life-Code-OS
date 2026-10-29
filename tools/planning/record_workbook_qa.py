"""Read-only exported-cell parity audit; append its receipt to workbook QA."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'docs/upgrade/2026-10-20'
OUT=P/'00_Doc_truoc'

def read(rel):return json.loads((P/rel).read_text())
def cell(v):
    if isinstance(v,list) and all(not isinstance(x,(dict,list)) for x in v):
        return '\n'.join(str(x).lower() if isinstance(x,bool) else str(x) for x in v) or None
    if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False,separators=(',',':'))
    return None if v=='' else v

def main():
    xlsx=OUT/'Life_Code_OS_Execution_Board.xlsx'
    wb=load_workbook(xlsx,read_only=True,data_only=False)
    tasks=read('08_Thuc_thi_Kinh_doanh/EXECUTION_BACKLOG.json')
    contracts={r['task_id']:r for r in read('08_Thuc_thi_Kinh_doanh/TASK_EXECUTION_CONTRACTS.json')}
    tables={}
    tables['Backlog']=[]
    for t in tasks:
        c=contracts[t['task_id']]
        tables['Backlog'].append([t[k] for k in ['task_id','phase','workstream','title','owner','status','effort_low_person_days','effort_high_person_days','dependencies','gates','deliverable','acceptance']]+[t.get('named_owner',t.get('assignee','CHƯA PHÂN CÔNG')),c['allowed_paths'],c['required_tests'],c['evidence_path'],t['reviewer'],c['decision_dependencies'],c['required_negative_tests'],c['forbidden_actions'],c['rollback_requirement'],c['production_permission'],c['evidence_required'],c['claim_status']])
    packets={r['decision_id']:r for r in read('08_Thuc_thi_Kinh_doanh/DECISION_EXECUTION_PACKETS.json')}
    tables['Decisions']=[]
    for r in read('08_Thuc_thi_Kinh_doanh/DECISION_REGISTER.json'):
        p=packets[r['decision_id']]
        recommendation=p['recommended_option_id']+'\n'+p['recommendation_rationale']+'\n'+'\n'.join(p['recommendation_tradeoffs'])
        options='\n\n'.join(o['option_id']+' '+o['label']+'\n'+o['scope']+'\n'+'\n'.join(o['tradeoffs']) for o in p['options'])
        tables['Decisions'].append([r['decision_id'],r['topic'],p['status'],recommendation,p['approver_roles'],options,p['constraints'],p['acceptance_receipt']['required_fields'],p['decision_due_event']])
    for name,rel,fields in [
        ('Gates','08_Thuc_thi_Kinh_doanh/RELEASE_GATES.json',['gate_id','name','owner','status','scope','evidence_uri','fpos_gate_ids']),
        ('Screens','05_Frontend_UI/SCREEN_REGISTRY.json',['screen_id','name_vi','target_route','principal','api','task_id','states_acceptance']),
        ('SourceMap','09_Doi_chieu_Bang_chung/SOURCE_TO_TARGET_MATRIX.json',['mapping_id','surface','source_path','action','target','task_id','acceptance']),
        ('FPOS','07_Bao_mat_Van_hanh/FPOS_GATE_MAPPING.json',['fpos_gate_id','title','custom_gate_ids','owner_role','evidence_required','status'])
    ]:tables[name]=[[r.get(k) for k in fields] for r in read(rel)]
    tables['Products']=[]
    for r in read('01_Chien_luoc_San_pham/PRODUCT_CATALOG.json'):
        price=r['price']
        price_text='Currency: '+str(price['currency'] if price['currency'] is not None else 'CHƯA CHỐT')+'\nAmount: '+str(price['amount'] if price['amount'] is not None else 'CHƯA CHỐT')+'\n'+price['status']
        tables['Products'].append([r[k] for k in ['product_id','family','name_vi','phase','status','input_contract','output_contract']]+[price_text])
    tables['Sources']=[[r['id'],r['title'],r['url'],r.get('access_date'),r.get('access_scope',r.get('scope_read')),r.get('limitations',r.get('limitation'))] for r in read('09_Doi_chieu_Bang_chung/SOURCE_REGISTRY.json')+read('09_Doi_chieu_Bang_chung/RESEARCH_UPDATE_2026-10-20.json')]
    for name,rows in [('Trace',read('08_Thuc_thi_Kinh_doanh/REQUIREMENT_TRACE_MATRIX.json')),('ReadCoverage',read('00_Doc_truoc/READ_COVERAGE.json')['files'])]:
        keys=list(dict.fromkeys(k for r in rows for k in r))
        tables[name]=[[r.get(k) for k in keys] for r in rows]
    failures=[];checked=0
    for name,rows in tables.items():
        width=len(rows[0])
        actual=list(wb[name].iter_rows(min_row=5,max_row=len(rows)+4,max_col=width,values_only=True))
        for i,(expected,found) in enumerate(zip(rows,actual),5):
            for j,(a,b) in enumerate(zip(expected,found),1):
                checked+=1
                if cell(a)!=b:failures.append(f'{name} row{i} col{j}: exported value differs')
        for row in wb[name].iter_rows(min_row=len(rows)+5,max_col=width,values_only=True):
            if any(v is not None for v in row):failures.append(name+': extra nonempty data row')
    wb.close()
    cached=load_workbook(xlsx,read_only=True,data_only=True)
    counts={address:cached['Dashboard'][address].value for address in ['A5','B5','C5','D5']}
    if counts!={'A5':156,'B5':0,'C5':0,'D5':0}:failures.append('Dashboard cached totals mismatch')
    cached.close()
    if failures:raise RuntimeError('\n'.join(failures[:30]))
    qa=json.loads((OUT/'WORKBOOK_QA.json').read_text())
    qa.update({'evidence_status':'VERIFIED_WORKTREE_ONLY','all_exported_registry_cell_checks':checked,'all_exported_registry_cell_mismatches':0,'registry_row_counts':{k:len(v) for k,v in tables.items()},'dashboard_cached_totals':counts,'visual_review':'22 preview panes independently inspected; 5 Backlog and 2 Products panes rechecked after padding/price readability fixes. This does not visually inspect every data row or native Excel UI.','xlsx_sha256':hashlib.sha256(xlsx.read_bytes()).hexdigest()})
    (OUT/'WORKBOOK_QA.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'exported_cells_checked':checked,'mismatches':0,'rows':qa['registry_row_counts'],'dashboard':counts},ensure_ascii=False))

if __name__=='__main__':main()
