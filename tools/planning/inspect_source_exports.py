"""Extract every page/paragraph/cell from supplied binary planning exports."""
import json
from pathlib import Path
from collections import Counter
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/upgrade/sources/2026-10-11/package/00_Doc_truoc'
OUT = ROOT / 'docs/upgrade/sources/2026-10-11/export-text'
OUT.mkdir(exist_ok=True)
pdf = PdfReader(SOURCE / 'Life_Code_OS_Master_Upgrade_Plan.pdf')
pages = [{'page': i + 1, 'text': p.extract_text()} for i, p in enumerate(pdf.pages)]
(OUT / 'pdf-pages.json').write_text(json.dumps(pages, ensure_ascii=False, indent=2) + '\n')
doc = Document(SOURCE / 'Life_Code_OS_Master_Upgrade_Plan.docx')
paragraphs = [p.text for p in doc.paragraphs]
tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
(OUT / 'docx-content.json').write_text(json.dumps({'paragraphs': paragraphs, 'tables': tables}, ensure_ascii=False, indent=2) + '\n')
book = load_workbook(SOURCE / 'Life_Code_OS_Execution_Board.xlsx', data_only=False)
cached = load_workbook(SOURCE / 'Life_Code_OS_Execution_Board.xlsx', data_only=True)
sheets = []
for sheet in book:
    cells = []
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                cells.append({'cell': cell.coordinate, 'value': cell.value,
                              'cached': cached[sheet.title][cell.coordinate].value,
                              'number_format': cell.number_format})
    sheets.append({'sheet': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column, 'cells': cells})
(OUT / 'xlsx-cells.json').write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str) + '\n')
summary = {'pdf_pages': len(pages), 'docx_paragraphs': len(paragraphs), 'docx_tables': len(tables),
           'xlsx_tabs': len(sheets), 'sheets': [{k:v for k,v in s.items() if k != 'cells'} for s in sheets],
           'xlsx_formulas': sum(isinstance(c['value'], str) and c['value'].startswith('=') for s in sheets for c in s['cells']),
           'pdf_replacement_characters': sum(p['text'].count('\ufffd') for p in pages)}
(OUT / 'SUMMARY.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(summary, ensure_ascii=False, indent=2))
