"""Compile the revised topics into Markdown and a reviewable Word handbook."""
import json
import re
import hashlib
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'docs/upgrade/2026-10-20'
SOURCE=ROOT/'docs/upgrade/sources/2026-10-11/package'
OUT=P/'00_Doc_truoc'

def plain(s):
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',s)
    s=s.replace('**','').replace('`','')
    return s.strip()

def table(doc, lines):
    parsed=[[plain(c) for c in l.strip().strip('|').split('|')] for l in lines]
    parsed=[row for row in parsed if not all(re.fullmatch(r':?-+:?',x.strip()) for x in row)]
    if not parsed:return
    width=max(map(len,parsed)); t=doc.add_table(rows=0,cols=width)
    t.autofit=False
    widths=[6.8/width]*width
    if width>=3:
        widths=[1.1]+[(6.8-1.1)/(width-1)]*(width-1)
    for i,row in enumerate(parsed):
        cells=t.add_row().cells
        for j,cell in enumerate(cells):
            cell.width=Inches(widths[j]);cell.text=row[j] if j<len(row) else ''
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F5' if i==0 else 'FFFFFF');tcPr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(9.5);run.bold=i==0
        if i==0:
            header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph()

def markdown(doc, text):
    lines=text.splitlines();i=0;code=False;sources=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if line.startswith('```'):
            code=not code;continue
        if not line or line=='---':continue
        if line.startswith('|') and not code:
            group=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                group.append(lines[i]);i+=1
            table(doc,group);continue
        h=re.match(r'^(#{1,6})\s+(.+)',line)
        if h and not code:
            if plain(h[2])=='Nguồn đối chiếu':sources=True
            if sources and len(h[1])==3:
                # A bibliographic identifier is literal evidence, not prose to
                # normalize (R20-01, WCAG 2.2, legal IDs and versions must survive).
                p=doc.add_paragraph(plain(h[2]))
                p.paragraph_format.keep_with_next=True
                p.paragraph_format.keep_together=True
                p.paragraph_format.space_before=Pt(10)
                for run in p.runs:run.bold=True
                continue
            level=min(len(h[1]),3)
            title=plain(h[2])
            title=re.sub(r'\b(20\d{2})-(\d{2})-(\d{2})\b', r'\3 tháng \2 năm \1', title)
            title=re.sub(r'(?<=\d)\s*[-–—]\s*(?=[A-Za-z]*\d)', ' đến ', title)
            title=re.sub(r'[^\w\sÀ-ỹ]',' ',title);title=re.sub(r'\s+',' ',title).strip()
            doc.add_heading(title,level=level);continue
        if code:
            p=doc.add_paragraph(line)
            p.paragraph_format.space_after=Pt(1)
            for run in p.runs:run.font.name='Courier New';run.font.size=Pt(9)
        elif re.match(r'^[-*]\s+',line):
            doc.add_paragraph(plain(line[2:]),style='List Bullet')
        elif re.match(r'^\d+\.\s+',line):
            # Preserve the author's local list numbering instead of Word continuing
            # one automatic counter across unrelated chapters and amendments.
            p=doc.add_paragraph(plain(line))
            p.paragraph_format.left_indent=Inches(0.25)
            p.paragraph_format.first_line_indent=Inches(-0.25)
        else:
            p=doc.add_paragraph(plain(line))
            p.paragraph_format.widow_control=True
            if sources:
                p.paragraph_format.keep_together=True
                if line.startswith(('https://','http://')):
                    p.paragraph_format.keep_with_next=True

def main():
    index=json.loads((SOURCE/'00_Doc_truoc/DOCUMENT_INDEX.json').read_text())
    old=(SOURCE/'00_Doc_truoc/Life_Code_OS_Master_Upgrade_Plan.md').read_text()
    topics=[r for r in index if r.get('master')]
    def position(row):
        first=(SOURCE/row['path']).read_text().splitlines()[0]
        n=old.find(first)
        return n if n>=0 else 1_000_000+int(row['id'][1:])
    topics.sort(key=position)
    front=(OUT/'AMENDMENT_2026-10-20.md').read_text()
    parts=['# Life Code OS Kế hoạch hợp nhất chuẩn bị phát triển\n\nPhiên bản 2026-10-20.r2. Release HOLD.\n\n',front,
           '\n# Mục lục chuyên đề\n\n'+'\n\n'.join(f"{r['id']} {r['title']} — `{r['path']}`" for r in topics)+'\n',
           '\n# Các chuyên đề hợp nhất\n\n']
    for r in topics:
        parts.append((P/r['path']).read_text())
    parts.append('\n# Tài liệu kỹ thuật chi tiết bổ sung\n\nCác hợp đồng công thức, ma trận phương pháp, demo fixtures, OpenAPI, migration, UI đa thiết bị, AI/report và task acceptance được tra theo INDEX.html và các README thư mục. File JSON/schema là artifact máy đọc của cùng phiên bản. Toàn bộ 149 tệp gốc và ZIP được giữ trong sources cùng checksum.\n')
    parts.append('\n# Nguồn đối chiếu\n')
    for r in json.loads((P/'09_Doi_chieu_Bang_chung/SOURCE_REGISTRY.json').read_text()):
        limitation=r.get('limitations') or 'Nguồn lịch sử chưa ghi; cần tái kiểm tra trước áp dụng'
        parts.append(f"\n### {r['id']} {r['title']}\n\n{r['url']}\n\nPhạm vi: {r.get('access_scope','')}. Giới hạn: {limitation}.\n")
    for r in json.loads((P/'09_Doi_chieu_Bang_chung/RESEARCH_UPDATE_2026-10-20.json').read_text()):
        parts.append(f"\n### {r['id']} {r['title']}\n\n{r['url']}\n\nĐọc 20/10: {r['scope_read']}. Giới hạn: {r['limitation']}.\n")
    combined='\n\n---\n\n'.join(parts)+'\n'
    (OUT/'Life_Code_OS_Master_Upgrade_Plan.md').write_text(combined)
    doc=Document()
    s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(0.7);s.bottom_margin=Inches(0.7);s.left_margin=Inches(0.85);s.right_margin=Inches(0.85)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(11 if name in ['Normal','List Bullet','List Number'] else 24 if name=='Title' else 16 if name=='Heading 1' else 13 if name=='Heading 2' else 11)
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
        if name.startswith('Heading'):st.paragraph_format.space_before=Pt(12);st.paragraph_format.keep_with_next=True
    doc.core_properties.title='Life Code OS Kế hoạch hợp nhất chuẩn bị phát triển'
    doc.core_properties.subject='Kế hoạch và hợp đồng chuẩn bị dev phiên bản 2026-10-20.r2'
    doc.core_properties.author='Life Code OS'
    doc.add_paragraph('Life Code OS',style='Title')
    doc.add_paragraph('Kế hoạch hợp nhất chuẩn bị phát triển',style='Subtitle')
    doc.add_paragraph('Phiên bản 2026-10-20.r2')
    doc.add_paragraph('Tài liệu dành cho Founder, Product, Method, Backend, Frontend, QA, Security và Operations. Bộ hồ sơ bảo toàn kế hoạch 11/10, bổ sung hợp đồng phát triển và demo tổng hợp. Release hiện HOLD cho tới khi có receipt đầy đủ của gate áp dụng.')
    doc.add_paragraph('Bắt đầu tại DEV_START_HERE_2026-10-20.md và INDEX.html. Các phương pháp chưa có source, rights và review không được phát hành như kết quả cá nhân. Chi tiết máy đọc nằm trong registry, schema, OpenAPI và các ma trận cùng gói.')
    doc.add_page_break()
    markdown(doc,'\n'.join(combined.splitlines()[1:]))
    footer=doc.sections[0].footer.paragraphs[0]
    footer.alignment=2
    run=footer.add_run('Life Code OS   2026-10-20.r2   ');run.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    output=OUT/'Life_Code_OS_Master_Upgrade_Plan.docx';doc.save(output)
    receipt={'scope':'MASTER_ASSEMBLY_NOT_RELEASE','version':'2026-10-20.r2','topics_preserved':len(topics),'markdown_sha256':hashlib.sha256(combined.encode()).hexdigest(),'docx_paragraphs':len(doc.paragraphs),'docx_tables':len(doc.tables),'source_files':[r['path'] for r in topics],'render_status':'NOT_CHECKED'}
    (OUT/'MASTER_ASSEMBLY.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in receipt.items() if k!='source_files'},ensure_ascii=False))

if __name__=='__main__':main()
